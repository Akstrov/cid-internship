"""
build_mapper5.py — SGOAM Mapping Tool v5

Key change: full nesting path used as key
  e.g. 0||Voie portée||Restriction d'exploitation||PTC
  e.g. 0||Voie Franchie||Type||Code

The path segments come directly from column positions in the Excel:
  c0 = level-0 section
  c1 = level-1 section
  c2 = level-2 section (or field if has type)
  type col = c3 for Franchissement/Ouvrage, c4 for Visite
"""

import re, json, sys
from collections import defaultdict


def extract_schema(sql_path):
    with open(sql_path, encoding="utf-8", errors="ignore") as f:
        content = f.read()
    tables_raw = re.findall(
        r"CREATE TABLE SGOAM\.(\w+)\s*\((.*?)\)\s*TABLESPACE", content, re.DOTALL
    )
    schema = {}
    for name, block in tables_raw:
        cols = {}
        for line in block.strip().split("\n"):
            line = line.strip().rstrip(",")
            if not line:
                continue
            if any(
                line.startswith(k)
                for k in ("CONSTRAINT", "PRIMARY", "FOREIGN", "UNIQUE", "CHECK")
            ):
                continue
            parts = line.split()
            if len(parts) >= 2 and re.match(r"^[A-Z][A-Z0-9_]*$", parts[0]):
                col_name = parts[0]
                type_parts = []
                for p in parts[1:]:
                    if p in ("NOT", "NULL", "DEFAULT", "ENABLE", "VALIDATE"):
                        break
                    type_parts.append(p)
                col_type = " ".join(type_parts).rstrip(",").replace(" BYTE)", ")")
                cols[col_name] = col_type
        schema[name] = cols
    return schema


def extract_fks(sql_path):
    with open(sql_path, encoding="utf-8", errors="ignore") as f:
        content = f.read()
    pattern = re.compile(
        r"ALTER TABLE SGOAM\.(\w+).*?FOREIGN KEY\s*\((.*?)\).*?REFERENCES SGOAM\.(\w+)\s*\((.*?)\)",
        re.DOTALL,
    )
    fks = defaultdict(list)
    for from_t, from_c, to_t, to_c in pattern.findall(content):
        from_c = ", ".join(c.strip() for c in from_c.split(","))
        to_c = ", ".join(c.strip() for c in to_c.split(","))
        fks[from_t].append({"to": to_t, "from_cols": from_c, "to_cols": to_c})
    return dict(fks)


def extract_pks(sql_path):
    with open(sql_path, encoding="utf-8", errors="ignore") as f:
        content = f.read()
    pattern = re.compile(
        r"CREATE UNIQUE INDEX SGOAM\.\w+ ON SGOAM\.(\w+)\s*\((.*?)\)", re.DOTALL
    )
    pks = {}
    for table, cols in pattern.findall(content):
        pks[table] = [c.strip() for c in cols.split(",")]
    return pks


def parse_sheet(ws, type_col, name_cols_outer_to_inner, onglet_col=None):
    """
    Parse a sheet into fields with full nesting path.
    name_cols_outer_to_inner: column indices ordered from outermost to innermost.
    e.g. (0, 1, 2) means c0=top-level section, c1=sub-section, c2=field or sub-sub-section.
    A row with a value but no type = section header at that depth level.
    A row with a type = data field. Its path = all section names at shallower levels.
    """
    SKIP_ENTIRELY = {
        'Données fiche "Données Franchissement"',
        'Fiche "Détails de l\'ouvrage" --> Niveau 1ère inspection spéciale',
        "(accessible à partir de la fiche de franchissement, "
        "qui peut comporter 1 ou plusieurs OA)",
        'Fenêtre visite (bouton "allez au niveau visite")',
    }
    NAV_BUTTONS = {
        "Historique",
        "Aide",
        "Photo de Franchissement",
        "Photos visite",
        "Allez au niveau visite",
        "Retour au niveau Franchissement",
        "Retour Ouvrage",
        "Retour Franchissement",
        "Saisir/Afficher Photos",
        "Ajouter un ouvrage",
    }

    rows = [list(r) for r in ws.iter_rows(values_only=True) if any(v for v in r)]
    n = len(name_cols_outer_to_inner)
    path = [None] * n
    fields = []

    for r in rows:
        while len(r) < 12:
            r.append(None)

        if (
            onglet_col is not None
            and r[onglet_col]
            and str(r[onglet_col]).startswith("Onglet")
        ):
            path[0] = str(r[onglet_col]).replace("Onglet ", "").strip('"').strip("'")
            for i in range(1, n):
                path[i] = None
            continue

        typ = str(r[type_col]).strip() if r[type_col] else ""

        # Find the deepest (innermost) column that has a value
        name = None
        depth = None
        for d in range(n - 1, -1, -1):
            ci = name_cols_outer_to_inner[d]
            if r[ci] is not None:
                name = str(r[ci]).strip()
                depth = d
                break

        if not name:
            continue
        if name in SKIP_ENTIRELY:
            continue

        if not typ:
            path[depth] = name
            for deeper in range(depth + 1, n):
                path[deeper] = None
            continue

        segments = [p for p in path[:depth] if p is not None]
        fields.append(
            {
                "path": segments,
                "field": name,
                "type": typ,
                "is_nav": name in NAV_BUTTONS,
            }
        )

    return fields


def parse_dashboard(xlsx_path):
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path)
    screens = []

    # Franchissement: type in c3, names in c2>c1>c0
    ws = wb["Existant - Franchissement"]
    fields = parse_sheet(ws, type_col=3, name_cols_outer_to_inner=(0, 1, 2))
    screens.append({"screen": "Franchissement", "fields": fields})

    # Détails Ouvrage: type in c3, names in c2>c1>c0, onglet in c0
    ws = wb["Existant - Détails Ouvrage"]
    fields = parse_sheet(
        ws, type_col=3, name_cols_outer_to_inner=(0, 1, 2), onglet_col=0
    )
    screens.append({"screen": "Détails Ouvrage", "fields": fields})

    # Visite: type in c4, names in c3>c2>c1>c0, onglet in c0
    ws = wb["Existant - Visite"]
    # Special handling for S/F/E/P unnamed rows
    SKIP_ENTIRELY = {
        'Données fiche "Données Franchissement"',
        'Fiche "Détails de l\'ouvrage" --> Niveau 1ère inspection spéciale',
        "(accessible à partir de la fiche de franchissement, "
        "qui peut comporter 1 ou plusieurs OA)",
        'Fenêtre visite (bouton "allez au niveau visite")',
    }
    NAV_BUTTONS = {
        "Historique",
        "Aide",
        "Photo de Franchissement",
        "Photos visite",
        "Allez au niveau visite",
        "Retour au niveau Franchissement",
        "Retour Ouvrage",
        "Retour Franchissement",
        "Saisir/Afficher Photos",
        "Ajouter un ouvrage",
    }
    sfep_labels = [
        "Note S - Superstructure",
        "Note F - Fondations et Appuis",
        "Note E - Equipements",
        "Note P - Protections et voisinage",
    ]
    rows_v = [list(r) for r in ws.iter_rows(values_only=True) if any(v for v in r)]
    path = [None, None, None, None]
    fields_v = []
    sfep_injected = False
    for r in rows_v:
        while len(r) < 12:
            r.append(None)
        if r[0] and str(r[0]).startswith("Onglet"):
            path[0] = str(r[0]).replace("Onglet ", "").strip('"').strip("'")
            path[1] = path[2] = path[3] = None
            continue
        typ = str(r[4]).strip() if r[4] else ""
        name = None
        col_depth = None
        for d in range(3, -1, -1):
            if r[d] is not None:
                name = str(r[d]).strip()
                col_depth = d
                break
        if not name:
            if typ == "Valeur automatique" and not sfep_injected:
                segs = [p for p in path if p is not None]
                for label in sfep_labels:
                    fields_v.append(
                        {
                            "path": segs[:],
                            "field": label,
                            "type": "Valeur automatique",
                            "is_nav": False,
                        }
                    )
                sfep_injected = True
            continue
        if name in SKIP_ENTIRELY:
            continue
        if not typ:
            path[col_depth] = name
            for deeper in range(col_depth + 1, 4):
                path[deeper] = None
            continue
        segs = [
            p for p in path[:col_depth] if p is not None
        ]  # col_depth is col index, path[i] is section at col i
        fields_v.append(
            {"path": segs, "field": name, "type": typ, "is_nav": name in NAV_BUTTONS}
        )
    screens.append({"screen": "Visite", "fields": fields_v})

    return screens


def mkey_py(screen_idx, path, field):
    parts = [str(screen_idx)] + path + [field]
    return "||".join(parts)


def generate_html(schema, fks, pks, screens, out_path):
    schema_js = json.dumps(schema, ensure_ascii=False)
    fks_js = json.dumps(fks, ensure_ascii=False)
    pks_js = json.dumps(pks, ensure_ascii=False)
    screens_js = json.dumps(screens, ensure_ascii=False)
    total = sum(len(s["fields"]) for s in screens)

    with open(out_path, "w", encoding="utf-8") as f:
        f.write(
            """<!DOCTYPE html>
<html lang="fr">
<head>
<meta charset="UTF-8">
<title>SGOAM Mapping Tool v5</title>
<style>
*{box-sizing:border-box;margin:0;padding:0}
body{font-family:system-ui,sans-serif;background:#f0f2f5;color:#1a1a1a;font-size:13px;height:100vh;display:flex;flex-direction:column;overflow:hidden}
header{background:#1F4E79;color:#fff;padding:0 16px;height:46px;display:flex;align-items:center;gap:12px;flex-shrink:0}
header h1{font-size:14px;font-weight:500;white-space:nowrap}
.pw{flex:1;min-width:100px;max-width:240px}
.pb{height:6px;background:rgba(255,255,255,.2);border-radius:3px;overflow:hidden;margin-top:3px}
.pf{height:100%;background:#70AD47;border-radius:3px;transition:width .3s}
.pl{font-size:11px;opacity:.75}
.hbtn{padding:5px 11px;font-size:12px;border:1px solid rgba(255,255,255,.35);background:rgba(255,255,255,.12);color:#fff;border-radius:5px;cursor:pointer;white-space:nowrap}
.hbtn:hover{background:rgba(255,255,255,.25)}
.hbtn.danger{background:rgba(220,50,50,.25)}
.stabs{display:flex;background:#fff;border-bottom:1px solid #ddd;flex-shrink:0}
.stab{padding:8px 18px;font-size:12px;font-weight:500;cursor:pointer;border:none;background:none;color:#888;border-bottom:2px solid transparent}
.stab.active{color:#1F4E79;border-bottom-color:#1F4E79}
.panels{display:grid;grid-template-columns:300px 1fr 1fr;flex:1;overflow:hidden}
.panel{display:flex;flex-direction:column;background:#fff;border-right:1px solid #e0e0e0;overflow:hidden}
.panel:last-child{border-right:none}
.ph{padding:7px 10px;font-size:11px;font-weight:600;text-transform:uppercase;letter-spacing:.05em;border-bottom:1px solid #eee;background:#fafafa;color:#555;flex-shrink:0;display:flex;align-items:center;justify-content:space-between}
.ph .cnt{font-size:10px;font-weight:400;color:#bbb;text-transform:none;letter-spacing:0}
.sb{padding:6px 8px;border-bottom:1px solid #eee;flex-shrink:0}
.sb input{width:100%;padding:5px 8px;font-size:12px;border:1px solid #ddd;border-radius:5px;outline:none}
.sb input:focus{border-color:#2E75B6}
.slist{overflow-y:auto;flex:1}
.sechdr{padding:5px 10px 2px;font-size:10px;font-weight:600;text-transform:uppercase;letter-spacing:.05em;color:#aaa;background:#f8f8f8;border-bottom:1px solid #f0f0f0;border-top:1px solid #f0f0f0;position:sticky;top:0;z-index:1}
.frow{padding:6px 10px;cursor:pointer;border-bottom:1px solid #f5f5f5;display:flex;align-items:flex-start;gap:7px}
.frow:hover{background:#f0f6ff}
.frow.active{background:#ddeeff}
.dot{width:7px;height:7px;border-radius:50%;flex-shrink:0;margin-top:4px}
.de{background:#ddd}
.dm{background:#70AD47}
.finfo{flex:1;min-width:0}
.fname{font-size:12px;line-height:1.35;word-break:break-word}
.fpath{font-size:10px;color:#bbb;margin-top:1px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
.ftype{font-size:10px;color:#bbb;white-space:nowrap}
.trow{padding:7px 10px;cursor:pointer;border-bottom:1px solid #f5f5f5;display:flex;align-items:center;justify-content:space-between;gap:6px}
.trow:hover{background:#f0f6ff}
.trow.active{background:#ddeeff;font-weight:500;color:#0C447C}
.tname{font-family:monospace;font-size:12px}
.tcnt{font-size:10px;color:#bbb;flex-shrink:0}
.trow.mm .tname{color:#bbb}
.crow{padding:7px 12px;cursor:pointer;border-bottom:1px solid #f5f5f5;display:flex;align-items:center;gap:6px}
.crow:hover{background:#f0fff4}
.crow.sel{background:#e2f0d9}
.cname{font-family:monospace;font-size:12px;flex:1}
.ctype{font-size:11px;color:#bbb}
.badge{font-size:10px;padding:1px 5px;border-radius:3px;white-space:nowrap;font-weight:500}
.fkb{background:#EBF3FB;color:#1F4E79}
.pkb{background:#FFF0E0;color:#8B4513}
.fkbox{border-top:1px solid #eee;background:#fffdf0;padding:10px 12px;font-size:12px;flex-shrink:0;max-height:120px;overflow-y:auto}
.fktitle{font-size:10px;font-weight:600;text-transform:uppercase;letter-spacing:.05em;color:#aaa;margin-bottom:5px}
.fkchain{font-family:monospace;font-size:11px;color:#555;line-height:1.9}
.mappings-panel{border-top:1px solid #eee;background:#f8fffe;flex-shrink:0;max-height:160px;overflow-y:auto}
.mp-header{padding:7px 12px;font-size:10px;font-weight:600;text-transform:uppercase;letter-spacing:.05em;color:#aaa;display:flex;align-items:center;justify-content:space-between;background:#f0faf8;border-bottom:1px solid #eee}
.mp-add{font-size:11px;color:#1F4E79;cursor:pointer;text-transform:none;letter-spacing:0;font-weight:500}
.mp-row{display:flex;align-items:center;gap:6px;padding:5px 12px;border-bottom:1px solid #f0f0f0;font-size:12px}
.mp-tbl{font-family:monospace;color:#1a1a1a;font-weight:500}
.mp-col{font-family:monospace;color:#375623}
.mp-del{margin-left:auto;color:#e74c3c;cursor:pointer;font-size:14px;line-height:1;flex-shrink:0;padding:0 2px}
.mp-empty{padding:10px 12px;font-size:12px;color:#ccc;font-style:italic}
.notebox{border-top:1px solid #eee;background:#fff;flex-shrink:0}
.notebox textarea{width:100%;padding:6px 12px 10px;font-size:12px;border:none;outline:none;resize:none;height:66px;font-family:system-ui;color:#444}
.notebox textarea::placeholder{color:#ccc}
.nblabel{padding:6px 12px 2px;font-size:10px;color:#aaa;font-weight:600;text-transform:uppercase;letter-spacing:.05em}
.mmrow{padding:5px 10px;border-bottom:1px solid #eee;font-size:11px;color:#888;display:flex;align-items:center;gap:5px;cursor:pointer;user-select:none;flex-shrink:0}
.empty{display:flex;align-items:center;justify-content:center;color:#ccc;font-size:13px;text-align:center;padding:20px;line-height:1.6;min-height:80px}
.sel-banner{background:#FFF3CD;border-bottom:1px solid #ffd;padding:6px 12px;font-size:11px;color:#856404;flex-shrink:0;display:none;align-items:center;justify-content:space-between}
.sel-banner.vis{display:flex}
</style>
</head>
<body>
<header>
  <h1>SGOAM — Mapping Dashboard → Base de données (v5)</h1>
  <div class="pw">
    <div class="pl" id="prog-label">0 / """
            + str(total)
            + """ mappés</div>
    <div class="pb"><div class="pf" id="prog-fill" style="width:0%"></div></div>
  </div>
  <button class="hbtn" onclick="saveProgress()">Sauvegarder</button>
  <button class="hbtn" onclick="loadProgress()">Charger</button>
  <button class="hbtn" onclick="exportBoth()">Exporter CSV + JSON</button>
  <button class="hbtn danger" onclick="clearAll()">Réinitialiser</button>
  <input type="file" id="load-input" accept=".json" style="display:none" onchange="onLoadFile(event)">
</header>

<div class="stabs" id="stabs"></div>

<div class="panels">
  <div class="panel">
    <div class="ph">Champs dashboard <span class="cnt" id="fc"></span></div>
    <div class="sb"><input type="text" id="fsrch" placeholder="Rechercher un champ..." oninput="renderFields()"></div>
    <div class="slist" id="flist"></div>
  </div>
  <div class="panel">
    <div class="ph">Tables <span class="cnt" id="tc"></span></div>
    <div class="sb"><input type="text" id="tsrch" placeholder="Rechercher une table..." oninput="renderTables()"></div>
    <div class="mmrow"><input type="checkbox" id="mmcb" checked onchange="renderTables()"><label for="mmcb">Tables _MM</label></div>
    <div class="slist" id="tlist"><div class="empty">Sélectionnez un champ</div></div>
  </div>
  <div class="panel">
    <div class="ph">Colonnes <span class="cnt" id="cc"></span></div>
    <div class="sb"><input type="text" id="csrch" placeholder="Rechercher une colonne..." oninput="renderCols()"></div>
    <div class="sel-banner" id="sel-banner">
      <span id="sel-banner-text"></span>
      <span style="cursor:pointer;font-weight:500" onclick="cancelExtra()">Annuler</span>
    </div>
    <div class="slist" id="clist"><div class="empty">Sélectionnez une table</div></div>
    <div class="fkbox" id="fkbox" style="display:none">
      <div class="fktitle">Contraintes FK de cette table</div>
      <div class="fkchain" id="fkchain"></div>
    </div>
    <div class="mappings-panel" id="mappings-panel" style="display:none">
      <div class="mp-header">
        Colonnes mappées
        <span class="mp-add" onclick="startAddExtra()">+ Ajouter une colonne</span>
      </div>
      <div id="mp-list"></div>
    </div>
    <div class="notebox" id="notebox" style="display:none">
      <div class="nblabel">Note / justification</div>
      <textarea id="note-input" placeholder="Pourquoi ce mapping ? Cas particulier ?" oninput="saveNote()"></textarea>
    </div>
  </div>
</div>

<script>
const SCHEMA="""
            + schema_js
            + """;
const FKS="""
            + fks_js
            + """;
const PKS="""
            + pks_js
            + """;
const SCREENS="""
            + screens_js
            + """;
const TOTAL="""
            + str(total)
            + """;
const ALL_TABLES=Object.keys(SCHEMA).sort();

let curScreen=0, selField=null, selTable=null, addingExtra=false;
let mappings=JSON.parse(localStorage.getItem('sgoam_map5')||'{}');

// Key = screenIdx || path[0] || path[1] || ... || field
function mkey(si, path, field){
  return [si, ...path, field].join('||');
}

function save(){localStorage.setItem('sgoam_map5',JSON.stringify(mappings));updateProg();}

function updateProg(){
  const d=Object.keys(mappings).filter(k=>mappings[k]&&mappings[k].cols&&mappings[k].cols.length>0).length;
  document.getElementById('prog-label').textContent=d+' / '+TOTAL+' mappés';
  document.getElementById('prog-fill').style.width=(d/TOTAL*100).toFixed(1)+'%';
}

function renderTabs(){
  document.getElementById('stabs').innerHTML=SCREENS.map((s,i)=>{
    const tot=s.fields.length;
    const done=s.fields.filter(f=>{const m=mappings[mkey(i,f.path,f.field)];return m&&m.cols&&m.cols.length>0;}).length;
    return `<button class="stab ${i===curScreen?'active':''}" onclick="switchScreen(${i})">${s.screen} <span style="font-size:10px;opacity:.6">${done}/${tot}</span></button>`;
  }).join('');
}

function switchScreen(i){
  curScreen=i; selField=null; selTable=null; addingExtra=false;
  renderTabs(); renderFields(); renderTables(); renderCols(); updateMappingsPanel(); updateNoteBox();
}

function renderFields(){
  const q=document.getElementById('fsrch').value.toLowerCase();
  // Group by top-level path segment (first element) or 'En-tête' if no path
  const groups={};
  let tot=0;
  SCREENS[curScreen].fields.forEach(f=>{
    if(q&&!f.field.toLowerCase().includes(q)&&!f.path.join(' ').toLowerCase().includes(q))return;
    const grp=f.path.length>0?f.path[0]:'En-tête';
    if(!groups[grp])groups[grp]=[];
    groups[grp].push(f); tot++;
  });
  document.getElementById('fc').textContent=tot+' champs';
  let html='';
  for(const[grp,flds]of Object.entries(groups)){
    html+=`<div class="sechdr">${grp}</div>`;
    flds.forEach(f=>{
      const k=mkey(curScreen,f.path,f.field);
      const m=mappings[k];
      const mapped=m&&m.cols&&m.cols.length>0;
      const act=selField&&mkey(curScreen,selField.path,selField.field)===k?'active':'';
      const esc=JSON.stringify(f).replace(/"/g,'&quot;');
      const navStyle=f.is_nav?'opacity:.55;font-style:italic':'';
      const navTag=f.is_nav?'<span style="font-size:9px;background:#ffeaa7;color:#856404;padding:1px 4px;border-radius:2px;margin-left:3px">nav</span>':'';
      // Show sub-path (everything after the first segment)
      const subPath=f.path.slice(1).join(' > ');
      const colCount=mapped?`<span style="font-size:10px;background:#e2f0d9;color:#375623;padding:1px 5px;border-radius:3px;white-space:nowrap">${m.cols.length} col${m.cols.length>1?'s':''}</span>`:'';
      html+=`<div class="frow ${act}" style="${navStyle}" onclick="selF(${esc})">
        <div class="dot ${mapped?'dm':'de'}"></div>
        <div class="finfo">
          <div class="fname">${f.field}${navTag}</div>
          ${subPath?`<div class="fpath">${subPath}</div>`:''}
        </div>
        ${colCount||`<div class="ftype">${f.type}</div>`}
      </div>`;
    });
  }
  document.getElementById('flist').innerHTML=html||'<div class="empty">Aucun résultat</div>';
}

function selF(f){
  selField=f; selTable=null; addingExtra=false;
  renderFields(); renderTables(); renderCols(); updateMappingsPanel(); updateNoteBox();
  setTimeout(()=>{const el=document.querySelector('.frow.active');if(el)el.scrollIntoView({block:'nearest'});},30);
}

function renderTables(){
  if(!selField){
    document.getElementById('tlist').innerHTML='<div class="empty">Sélectionnez un champ</div>';
    document.getElementById('tc').textContent=''; return;
  }
  const q=document.getElementById('tsrch').value.toLowerCase();
  const showMM=document.getElementById('mmcb').checked;
  const tables=ALL_TABLES.filter(t=>{
    if(!showMM&&t.endsWith('_MM'))return false;
    if(q&&!t.toLowerCase().includes(q))return false;
    return true;
  });
  document.getElementById('tc').textContent=tables.length+' tables';
  const html=tables.map(t=>{
    const nc=Object.keys(SCHEMA[t]||{}).length;
    const mm=t.endsWith('_MM');
    const act=selTable===t?'active':'';
    return `<div class="trow ${mm?'mm':''} ${act}" onclick="selT('${t}')">
      <span class="tname">${t}</span><span class="tcnt">${nc} cols</span>
    </div>`;
  }).join('');
  document.getElementById('tlist').innerHTML=html||'<div class="empty">Aucun résultat</div>';
}

function selT(t){selTable=t; renderTables(); renderCols(); updateFKBox(t);}

function renderCols(){
  const banner=document.getElementById('sel-banner');
  if(addingExtra&&selField){
    banner.classList.add('vis');
    document.getElementById('sel-banner-text').textContent='Mode ajout — cliquez une colonne pour l\\'ajouter à "'+selField.field+'"';
  } else banner.classList.remove('vis');

  if(!selTable){
    document.getElementById('clist').innerHTML='<div class="empty">Sélectionnez une table</div>';
    document.getElementById('cc').textContent='';
    document.getElementById('fkbox').style.display='none'; return;
  }
  const q=document.getElementById('csrch').value.toLowerCase();
  const cols=SCHEMA[selTable]||{};
  const fksOut=(FKS[selTable]||[]).map(r=>r.from_cols).join(' ');
  const pkCols=PKS[selTable]||[];
  const filtered=Object.entries(cols).filter(([c])=>!q||c.toLowerCase().includes(q));
  document.getElementById('cc').textContent=filtered.length+' colonnes';
  const k=selField?mkey(curScreen,selField.path,selField.field):null;
  const cur=k?(mappings[k]||{}):{}; 
  const assignedCols=(cur.cols||[]).filter(c=>c.table===selTable).map(c=>c.col);
  const html=filtered.map(([cn,ct])=>{
    const hasFk=fksOut.includes(cn);
    const isPk=pkCols.includes(cn);
    const isSel=assignedCols.includes(cn);
    const cnEsc=cn.replace(/\\\\/g,'\\\\\\\\').replace(/'/g,"\\\\'");
    const addStyle=addingExtra?'style="background:#fffbe6"':'';
    return `<div class="crow ${isSel?'sel':''}" ${addStyle} onclick="handleColClick('${selTable}','${cnEsc}')">
      <span class="cname">${cn}</span>
      <span class="ctype">${ct}</span>
      ${isPk?'<span class="badge pkb">PK</span>':''}
      ${hasFk?'<span class="badge fkb">FK</span>':''}
    </div>`;
  }).join('');
  document.getElementById('clist').innerHTML=html||'<div class="empty">Aucun résultat</div>';
}

function updateFKBox(t){
  const box=document.getElementById('fkbox');
  const fksOut=FKS[t]||[];
  if(!fksOut.length){box.style.display='none';return;}
  box.style.display='block';
  document.getElementById('fkchain').innerHTML=fksOut.map(r=>
    `<span style="color:#1a1a1a">${t}.${r.from_cols}</span> <span style="color:#bbb">→</span> ${r.to}.${r.to_cols}`
  ).join('<br>');
}

function handleColClick(table,col){
  if(!selField)return;
  if(addingExtra) addColToMapping(table,col);
  else assignFirstCol(table,col);
}

function assignFirstCol(table,col){
  const k=mkey(curScreen,selField.path,selField.field);
  const existing=mappings[k]||{};
  mappings[k]={
    screen:SCREENS[curScreen].screen,
    path:selField.path,
    field:selField.field,
    type:selField.type,
    cols:[{table,col}],
    note:existing.note||''
  };
  save(); renderFields(); renderCols(); updateMappingsPanel(); autoAdvance();
}

function addColToMapping(table,col){
  const k=mkey(curScreen,selField.path,selField.field);
  if(!mappings[k]) mappings[k]={screen:SCREENS[curScreen].screen,path:selField.path,field:selField.field,type:selField.type,cols:[],note:''};
  const already=mappings[k].cols.some(c=>c.table===table&&c.col===col);
  if(!already) mappings[k].cols.push({table,col});
  save(); renderFields(); renderCols(); updateMappingsPanel();
}

function removeCol(table,col){
  if(!selField)return;
  const k=mkey(curScreen,selField.path,selField.field);
  if(!mappings[k])return;
  mappings[k].cols=mappings[k].cols.filter(c=>!(c.table===table&&c.col===col));
  save(); renderFields(); renderCols(); updateMappingsPanel();
}

function startAddExtra(){if(!selField)return; addingExtra=true; renderCols();}
function cancelExtra(){addingExtra=false; renderCols();}

function updateMappingsPanel(){
  const panel=document.getElementById('mappings-panel');
  const list=document.getElementById('mp-list');
  if(!selField){panel.style.display='none';return;}
  panel.style.display='block';
  const k=mkey(curScreen,selField.path,selField.field);
  const m=mappings[k];
  const cols=(m&&m.cols)||[];
  if(!cols.length){
    list.innerHTML='<div class="mp-empty">Aucune colonne mappée — cliquez sur une colonne ci-dessus</div>';
    return;
  }
  list.innerHTML=cols.map(c=>{
    const tEsc=c.table.replace(/'/g,"\\'");
    const cEsc=c.col.replace(/'/g,"\\'");
    return `<div class="mp-row">
      <span class="mp-tbl">${c.table}</span>
      <span style="color:#aaa">.</span>
      <span class="mp-col">${c.col}</span>
      <span class="mp-del" onclick="removeCol('${tEsc}','${cEsc}')" title="Supprimer">×</span>
    </div>`;
  }).join('');
}

function updateNoteBox(){
  const nb=document.getElementById('notebox');
  const ni=document.getElementById('note-input');
  if(!selField){nb.style.display='none';return;}
  nb.style.display='block';
  const k=mkey(curScreen,selField.path,selField.field);
  const m=mappings[k];
  ni.value=m?m.note||'':'';
}

function saveNote(){
  if(!selField)return;
  const k=mkey(curScreen,selField.path,selField.field);
  if(!mappings[k]) mappings[k]={screen:SCREENS[curScreen].screen,path:selField.path,field:selField.field,type:selField.type,cols:[],note:''};
  mappings[k].note=document.getElementById('note-input').value;
  save();
}

function autoAdvance(){
  const fields=SCREENS[curScreen].fields;
  if(!selField)return;
  const curKey=mkey(curScreen,selField.path,selField.field);
  const idx=fields.findIndex(f=>mkey(curScreen,f.path,f.field)===curKey);
  for(let i=idx+1;i<fields.length;i++){
    const m=mappings[mkey(curScreen,fields[i].path,fields[i].field)];
    if(!m||!m.cols||!m.cols.length){selF(fields[i]);return;}
  }
}

function saveProgress(){
  const a=document.createElement('a');
  a.href='data:application/json;charset=utf-8,'+encodeURIComponent(JSON.stringify({version:5,mappings},null,2));
  a.download='sgoam_progress_v5.json'; a.click();
}

function loadProgress(){document.getElementById('load-input').click();}

function onLoadFile(e){
  const file=e.target.files[0]; if(!file)return;
  const reader=new FileReader();
  reader.onload=ev=>{
    try{
      const parsed=JSON.parse(ev.target.result);
      mappings=parsed.mappings||parsed;
      save(); renderTabs(); renderFields(); renderTables(); renderCols();
      updateMappingsPanel(); updateNoteBox();
      alert('Chargé — '+Object.keys(mappings).length+' entrées.');
    }catch(err){alert('Erreur: '+err.message);}
  };
  reader.readAsText(file); e.target.value='';
}

function exportBoth(){
  const rows=[['Ecran','Chemin','Champ','Type saisie','Table DB','Colonne DB','Note']];
  SCREENS.forEach((s,si)=>{
    s.fields.forEach(f=>{
      const m=mappings[mkey(si,f.path,f.field)]||{};
      const cols=m.cols||[];
      const pathStr=f.path.join(' > ');
      if(!cols.length){
        rows.push([s.screen,pathStr,f.field,f.type,'','',m.note||'']);
      } else {
        cols.forEach((c,ci)=>{
          rows.push([s.screen,pathStr,f.field,f.type,c.table,c.col,ci===0?m.note||'':'']);
        });
      }
    });
  });
  const csv=rows.map(r=>r.map(c=>'"'+String(c).replace(/"/g,'""')+'"').join(',')).join('\\n');
  dl('data:text/csv;charset=utf-8,'+encodeURIComponent('\\uFEFF'+csv),'sgoam_mapping.csv');
  const output={};
  SCREENS.forEach((s,si)=>{
    output[s.screen]=s.fields.map(f=>{
      const m=mappings[mkey(si,f.path,f.field)]||{};
      return {path:f.path,field:f.field,type:f.type,cols:m.cols||[],note:m.note||null};
    });
  });
  dl('data:application/json;charset=utf-8,'+encodeURIComponent(JSON.stringify(output,null,2)),'sgoam_mapping.json');
}

function dl(href,name){const a=document.createElement('a');a.href=href;a.download=name;a.click();}

function clearAll(){
  if(!confirm('Effacer tous les mappings ?'))return;
  mappings={}; save(); selField=null; selTable=null; addingExtra=false;
  renderTabs(); renderFields(); renderTables(); renderCols(); updateMappingsPanel(); updateNoteBox();
}

renderTabs(); renderFields(); updateProg();
</script>
</body>
</html>"""
        )

    print(f"Generated : {out_path}")
    print(f"  Screens : {len(screens)}")
    for s in screens:
        print(f"    {s['screen']}: {len(s['fields'])} fields")
    print(f"  Tables  : {len(schema)}")
    print(f"  PKs     : {len(pks)}")
    print(f"  FKs     : {sum(len(v) for v in fks.values())}")


if __name__ == "__main__":
    sql_path = sys.argv[1] if len(sys.argv) > 1 else "SGOAM2_sql.txt"
    xlsx_path = (
        sys.argv[2]
        if len(sys.argv) > 2
        else "Arborescences fenêtres SGOAM - Outil existant.xlsx"
    )
    out_path = sys.argv[3] if len(sys.argv) > 3 else "mapper.html"

    print("Extracting schema...")
    schema = extract_schema(sql_path)
    print("Extracting FK constraints...")
    fks = extract_fks(sql_path)
    print("Extracting PK indexes...")
    pks = extract_pks(sql_path)
    print("Parsing dashboard fields...")
    screens = parse_dashboard(xlsx_path)
    print("Generating HTML...")
    generate_html(schema, fks, pks, screens, out_path)
